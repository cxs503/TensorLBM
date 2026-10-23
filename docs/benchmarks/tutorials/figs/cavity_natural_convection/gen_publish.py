"""把本教程 Markdown 转成 Word / PDF 变体（供 tutorials/files/ 使用）。

只做本仓库教程用到的 Markdown 子集：标题、段落、引用、代码围栏、管道表格、
图片、分隔线。中文字体用系统的 DroidSansFallback（TrueType，reportlab 可嵌入）。
"""

from __future__ import annotations

import re
from pathlib import Path

TUT = Path(__file__).resolve().parents[2]          # docs/benchmarks/tutorials
MD = TUT / "cavity_natural_convection.md"
FILES = TUT / "files"
CJK_TTF = Path("/usr/share/fonts/google-droid-sans-fonts/DroidSansFallbackFull.ttf")
LAT_TTF = Path("/usr/share/fonts/dejavu-sans-fonts/DejaVuSans.ttf")
MONO_TTF = Path("/usr/share/fonts/dejavu-sans-fonts/DejaVuSansMono.ttf")


def _mono_font() -> Path:
    """等宽字体：优先系统 DejaVu Sans Mono，否则用 matplotlib 自带的同名 TTF。"""
    if MONO_TTF.exists():
        return MONO_TTF
    import matplotlib
    p = Path(matplotlib.__file__).parent / "mpl-data" / "fonts" / "ttf" / "DejaVuSansMono.ttf"
    return p


def cjk_runs(text: str) -> list[tuple[str, str]]:
    """把文本切成 (字体类别, 片段) 序列：'cjk' 走 Droid Sans Fallback，'lat' 走 DejaVu Sans。

    Droid Sans Fallback 只含 CJK 字形（无 ASCII/希腊字母/箭头），DejaVu Sans 无汉字；
    reportlab 不做字体回退，故按字符集切段并在 Paragraph 标记里交替指定字体。
    """

    def is_cjk(ch: str) -> bool:
        o = ord(ch)
        return (
            0x3000 <= o <= 0x303F          # CJK 标点
            or 0x3400 <= o <= 0x4DBF       # 扩展 A
            or 0x4E00 <= o <= 0x9FFF       # 基本汉字
            or 0xF900 <= o <= 0xFAFF       # 兼容汉字
            or 0xFF00 <= o <= 0xFFEF       # 全角
        )

    runs: list[tuple[str, str]] = []
    for ch in text:
        kind = "cjk" if is_cjk(ch) else "lat"
        if runs and runs[-1][0] == kind:
            runs[-1] = (kind, runs[-1][1] + ch)
        else:
            runs.append((kind, ch))
    return runs


def markup(text: str) -> str:
    """把纯文本转成带字体切换的 reportlab Paragraph 标记（含 XML 转义）。"""
    out = []
    for kind, seg in cjk_runs(text):
        esc = seg.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        name = "CJK" if kind == "cjk" else "LAT"
        out.append(f'<font name="{name}">{esc}</font>')
    return "".join(out)

INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")


def clean(text: str) -> str:
    """去掉行内 markdown 记号，保留文字（保留 `data/result_n*.json` 这类通配星号）。"""
    text = text.replace("\\|", "|")
    text = re.sub(r"!\[[^\]]*\]\([^)]*\)", "", text)
    text = re.sub(r"\[([^\]]*)\]\([^)]*\)", r"\1", text)
    text = re.sub(r"`([^`]*)`", r"\1", text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    if text.startswith("**") and text.endswith("**") and len(text) > 4:
        text = text[2:-2]
    text = re.sub(r"(?<!\*)\*([^*]+?)\*(?!\*)", r"\1", text)
    return text.strip()


def parse(md: str) -> list[dict]:
    """返回块列表；块类型：h1/h2/h3/quote/p/code/table/img/hr。"""
    blocks: list[dict] = []
    lines = md.split("\n")
    i = 0
    while i < len(lines):
        ln = lines[i]
        if not ln.strip():
            i += 1
            continue
        if ln.startswith("```"):
            lang = ln[3:].strip()
            body = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                body.append(lines[i])
                i += 1
            i += 1
            blocks.append({"t": "code", "lines": body, "lang": lang})
            continue
        m = re.match(r"^(#{1,3})\s+(.*)$", ln)
        if m:
            blocks.append({"t": f"h{len(m.group(1))}", "text": clean(m.group(2))})
            i += 1
            continue
        if ln.startswith("> "):
            blocks.append({"t": "quote", "text": clean(ln[2:])})
            i += 1
            continue
        m = re.match(r"^!\[([^\]]*)\]\(([^)]*)\)", ln)
        if m:
            blocks.append({"t": "img", "alt": clean(m.group(1)), "src": m.group(2)})
            i += 1
            continue
        if re.match(r"^(-{3,}|\*{3,})$", ln.strip()):
            blocks.append({"t": "hr"})
            i += 1
            continue
        if ln.lstrip().startswith("|"):
            rows = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                raw = lines[i].strip().strip("|")
                cells = [c.strip() for c in re.split(r"(?<!\\)\|", raw)]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append([clean(c) for c in cells])
                i += 1
            blocks.append({"t": "table", "rows": rows})
            continue
        para = [ln]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(
            r"^(#{1,3}\s|```|>\s|!\[|\||-{3,}$)", lines[i]
        ):
            para.append(lines[i])
            i += 1
        blocks.append({"t": "p", "text": clean(" ".join(para))})
    return blocks


# ---------- Word ----------

def to_docx(blocks: list[dict], out: Path) -> None:
    from docx import Document
    from docx.shared import Emu, Pt

    doc = Document()
    width = Emu(5_800_000)  # ≈16 cm 正文宽度
    for b in blocks:
        t = b["t"]
        if t == "h1":
            doc.add_heading(b["text"], level=1)
        elif t == "h2":
            doc.add_heading(b["text"], level=2)
        elif t == "h3":
            doc.add_heading(b["text"], level=3)
        elif t == "quote":
            p = doc.add_paragraph(b["text"])
            p.runs[0].italic = True
        elif t == "p":
            doc.add_paragraph(b["text"])
        elif t == "code":
            p = doc.add_paragraph()
            run = p.add_run("\n".join(b["lines"]))
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        elif t == "table":
            rows = b["rows"]
            if not rows:
                continue
            ncol = max(len(r) for r in rows)
            tab = doc.add_table(rows=0, cols=ncol)
            tab.style = "Table Grid"
            for r in rows:
                cells = tab.add_row().cells
                for k, val in enumerate(r):
                    if k < ncol:
                        cells[k].text = val
        elif t == "img":
            src = TUT / b["src"]
            if src.exists():
                doc.add_picture(str(src), width=width)
                cap = doc.add_paragraph(b["alt"])
                cap.runs[0].font.size = Pt(8.5)
        elif t == "hr":
            doc.add_paragraph()
    doc.save(out)


# ---------- PDF ----------

def to_pdf(blocks: list[dict], out: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_JUSTIFY
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (
        Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
    )
    from reportlab.lib.utils import ImageReader

    pdfmetrics.registerFont(TTFont("CJK", str(CJK_TTF)))
    pdfmetrics.registerFont(TTFont("LAT", str(LAT_TTF)))
    pdfmetrics.registerFont(TTFont("MONO", str(_mono_font())))
    base = ParagraphStyle("b", fontName="LAT", fontSize=9.5, leading=14, alignment=TA_JUSTIFY)
    styles = {
        "h1": ParagraphStyle("h1", parent=base, fontSize=16, leading=22, spaceBefore=6, spaceAfter=10),
        "h2": ParagraphStyle("h2", parent=base, fontSize=12.5, leading=18, spaceBefore=12, spaceAfter=6),
        "h3": ParagraphStyle("h3", parent=base, fontSize=11, leading=16, spaceBefore=8, spaceAfter=4),
        "p": base,
        "quote": ParagraphStyle("q", parent=base, textColor=colors.HexColor("#444444"),
                                leftIndent=10, spaceAfter=8),
        "code": ParagraphStyle("c", parent=base, fontSize=8, leading=11, fontName="MONO",
                               backColor=colors.HexColor("#f3f3f3"),
                               leftIndent=8, borderPadding=4),
        "cap": ParagraphStyle("cap", parent=base, fontSize=8.5, leading=12,
                              textColor=colors.HexColor("#333333"), spaceAfter=8),
    }
    story: list = []
    avail = A4[0] - 3.4 * cm
    for b in blocks:
        t = b["t"]
        if t in ("h1", "h2", "h3"):
            story.append(Paragraph(markup(b["text"]), styles[t]))
        elif t == "quote":
            story.append(Paragraph(markup(b["text"]), styles["quote"]))
        elif t == "p":
            story.append(Paragraph(markup(b["text"]), styles["p"]))
        elif t == "code":
            body = "<br/>".join(
                "".join(
                    f'<font name="{"CJK" if k == "cjk" else "MONO"}">{s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")}</font>'
                    for k, s in cjk_runs(ln)
                )
                for ln in b["lines"]
            )
            story.append(Paragraph(body, styles["code"]))
            story.append(Spacer(1, 4))
        elif t == "table":
            rows = b["rows"]
            if not rows:
                continue
            data = [[Paragraph(markup(str(c)), ParagraphStyle(
                "tc", parent=base, fontSize=7.4, leading=10)) for c in r] for r in rows]
            ncol = max(len(r) for r in rows)
            data = [r + [""] * (ncol - len(r)) for r in data]
            tab = Table(data, colWidths=[avail / ncol] * ncol, repeatRows=1)
            tab.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#999999")),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eaeaea")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]))
            story.append(tab)
            story.append(Spacer(1, 8))
        elif t == "img":
            src = TUT / b["src"]
            if src.exists():
                iw, ih = ImageReader(str(src)).getSize()
                w = avail
                story.append(Image(str(src), width=w, height=w * ih / iw))
            story.append(Paragraph(markup(b["alt"]), styles["cap"]))
        elif t == "hr":
            story.append(Spacer(1, 6))
    SimpleDocTemplate(str(out), pagesize=A4, leftMargin=1.7 * cm, rightMargin=1.7 * cm,
                      topMargin=1.6 * cm, bottomMargin=1.6 * cm,
                      title="cavity_natural_convection tutorial").build(story)


def main() -> None:
    FILES.mkdir(parents=True, exist_ok=True)
    blocks = parse(MD.read_text(encoding="utf-8"))
    to_docx(blocks, FILES / "cavity_natural_convection.docx")
    to_pdf(blocks, FILES / "cavity_natural_convection.pdf")
    print(f"wrote {FILES/'cavity_natural_convection.docx'} / .pdf ({len(blocks)} blocks)")


if __name__ == "__main__":
    main()